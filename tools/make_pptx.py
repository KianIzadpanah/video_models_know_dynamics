#!/usr/bin/env python3
"""
Build a PowerPoint deck for an experiment from its slides.json.

    python tools/make_pptx.py exp6

Writes ``<exp>-slides.pptx`` next to the experiment folders — the same place the
HTML deck goes, and outside the site repo, so it is never published.

Why this exists separately from the HTML deck: PowerPoint cannot play the HTML
deck's synchronised video grids, so the motion has to become *frames*. Every clip
is sampled at four fixed times and laid out as conditions-down, time-across, which
is the one arrangement that shows both what changed with K and what happened over
the six seconds. Nothing on a slide is a video, an animation or an embedded
object: it is pictures and text boxes, so it opens anywhere and prints.

Requires: python-pptx, Pillow, and ffmpeg on PATH.
"""

from __future__ import annotations

import argparse
import json
import re
import subprocess
import sys
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parent.parent          # …/site
OUT_DIR = ROOT.parent                                  # …/Experiments
CACHE = ROOT / ".pptx-frames"                          # extracted frames, gitignored

# ---------------------------------------------------------------------------
# light palette — deliberately plain: white ground, one accent, two verdict tints
# ---------------------------------------------------------------------------
INK = RGBColor(0x1A, 0x1D, 0x23)
INK2 = RGBColor(0x3D, 0x44, 0x4F)
MUTED = RGBColor(0x6B, 0x72, 0x80)
FAINT = RGBColor(0x9A, 0xA2, 0xAE)
ACCENT = RGBColor(0x1F, 0x5B, 0xD7)
KEY = RGBColor(0x0D, 0x6B, 0x4F)
HEAVY = RGBColor(0xB0, 0x43, 0x1D)
LINE = RGBColor(0xE3, 0xE6, 0xEA)
PANEL = RGBColor(0xF5, 0xF6, 0xF8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

RGB_INK = (0x1A, 0x1D, 0x23)
RGB_MUTED = (0x6B, 0x72, 0x80)
RGB_FAINT = (0x9A, 0xA2, 0xAE)
RGB_ACCENT = (0x1F, 0x5B, 0xD7)
RGB_LINE = (0xE3, 0xE6, 0xEA)
RGB_WHITE = (0xFF, 0xFF, 0xFF)

SANS = "Segoe UI"
MONO = "Consolas"
F_SANS = "C:/Windows/Fonts/segoeui.ttf"
F_BOLD = "C:/Windows/Fonts/segoeuib.ttf"
F_MONO = "C:/Windows/Fonts/consola.ttf"

W_IN, H_IN = 13.333, 7.5
MARGIN = 0.45

# four fixed sample times, so every clip is read on the same clock
SAMPLES = [0, 56, 112, 168]
TILE_PX = 320                      # per-tile width in the composites (~230 dpi)


def font(path: str, size: int):
    try:
        return ImageFont.truetype(path, size)
    except OSError:
        return ImageFont.load_default()


# ---------------------------------------------------------------------------
# frame extraction
# ---------------------------------------------------------------------------
def extract(src: Path, frames: list[int], tag: str) -> list[Path]:
    """Pull the requested frame indices out of one video, cached on disk."""
    outs = [CACHE / f"{tag}_f{n}.png" for n in frames]
    if all(o.is_file() for o in outs):
        return outs
    CACHE.mkdir(parents=True, exist_ok=True)
    sel = "+".join(f"eq(n\\,{n})" for n in frames)
    pat = CACHE / f"{tag}_seq%02d.png"
    cmd = ["ffmpeg", "-y", "-v", "error", "-i", str(src),
           "-vf", f"select='{sel}'", "-vsync", "0", "-frames:v", str(len(frames)),
           str(pat)]
    r = subprocess.run(cmd, capture_output=True, text=True)
    if r.returncode != 0:
        print(f"    ! ffmpeg failed on {src.name}: {(r.stderr or '').strip()[:120]}")
        return []
    seq = sorted(CACHE.glob(f"{tag}_seq*.png"))
    for i, o in enumerate(outs):
        if i < len(seq):
            seq[i].replace(o)
    for leftover in CACHE.glob(f"{tag}_seq*.png"):
        leftover.unlink()
    return [o for o in outs if o.is_file()]


# ---------------------------------------------------------------------------
# composites
# ---------------------------------------------------------------------------
def grid_image(rows: list[list[Path]], out: Path, tile_w: int = TILE_PX,
               gap: int = 7) -> Path | None:
    """rows[r][c] -> one image, white gaps. Missing cells are left blank."""
    if not rows or not rows[0]:
        return None
    probe = next((p for r in rows for p in r if p and p.is_file()), None)
    if probe is None:
        return None
    with Image.open(probe) as im:
        ar = im.width / im.height
    tile_h = int(round(tile_w / ar))
    ncol = max(len(r) for r in rows)
    W = ncol * tile_w + (ncol - 1) * gap
    H = len(rows) * tile_h + (len(rows) - 1) * gap
    canvas = Image.new("RGB", (W, H), RGB_WHITE)
    for ri, row in enumerate(rows):
        for ci, p in enumerate(row):
            if not p or not p.is_file():
                continue
            with Image.open(p) as im:
                canvas.paste(im.convert("RGB").resize((tile_w, tile_h), Image.LANCZOS),
                             (ci * (tile_w + gap), ri * (tile_h + gap)))
    out.parent.mkdir(parents=True, exist_ok=True)
    canvas.save(out, "JPEG", quality=88, optimize=True)
    return out


def pipeline_image(stages: list, side: list | None, out: Path) -> Path:
    """The pipeline, drawn light. Boxes come from slides.json, so there is no
    diagram asset to keep in sync with the text."""
    S = 3                                    # supersample, then downscale
    BW, BH, GAP = 250 * S, 108 * S, 52 * S
    n = len(stages)
    top = (150 * S) if side else (20 * S)
    W = n * BW + (n - 1) * GAP
    H = top + BH + 20 * S
    im = Image.new("RGB", (W, H), RGB_WHITE)
    d = ImageDraw.Draw(im)
    fb, fs = font(F_BOLD, 25 * S), font(F_SANS, 20 * S)

    def box(x, y, w, h, fill, outline):
        d.rounded_rectangle([x, y, x + w, y + h], radius=12 * S, fill=fill,
                            outline=outline, width=2 * S)

    def centre(txt, cx, cy, f, col):
        l, t, r, b = d.textbbox((0, 0), txt, font=f)
        d.text((cx - (r - l) / 2, cy - (b - t) / 2), txt, font=f, fill=col)

    def arrow(x1, y1, x2, y2):
        d.line([x1, y1, x2, y2], fill=RGB_FAINT, width=3 * S)
        if y1 == y2:
            d.polygon([(x2, y2), (x2 - 11 * S, y2 - 7 * S), (x2 - 11 * S, y2 + 7 * S)],
                      fill=RGB_FAINT)
        else:
            d.polygon([(x2, y2), (x2 - 7 * S, y2 - 11 * S), (x2 + 7 * S, y2 - 11 * S)],
                      fill=RGB_FAINT)

    for i, st in enumerate(stages):
        label, sub = (list(st) + ["", ""])[:2] if isinstance(st, (list, tuple)) else (st, "")
        x = i * (BW + GAP)
        box(x, top, BW, BH, (0xF5, 0xF6, 0xF8), RGB_LINE)
        centre(label, x + BW / 2, top + BH * 0.36, fb, RGB_INK)
        if sub:
            centre(sub, x + BW / 2, top + BH * 0.68, fs, RGB_MUTED)
        if i < n - 1:
            arrow(x + BW + 10 * S, top + BH / 2, x + BW + GAP - 10 * S, top + BH / 2)

    if side:
        cx = 2 * (BW + GAP) + BW / 2
        sw, sh = 420 * S, 88 * S
        box(cx - sw / 2, 18 * S, sw, sh, (0xE9, 0xF0, 0xFF), RGB_ACCENT)
        centre(side[0], cx, 18 * S + sh * 0.34, fb, RGB_ACCENT)
        if len(side) > 1:
            centre(side[1], cx, 18 * S + sh * 0.70, fs, RGB_MUTED)
        arrow(cx, 18 * S + sh + 8 * S, cx, top - 12 * S)

    im = im.resize((W // S, H // S), Image.LANCZOS)
    out.parent.mkdir(parents=True, exist_ok=True)
    im.save(out, "PNG")
    return out


def timelines_image(conds: list, total: int, out: Path) -> Path:
    """Ticks on a 0..total axis: one row per condition, so K reads as a picture."""
    S = 3
    LW, RH, GAP = 1500 * S, 46 * S, 26 * S
    LAB, NUM = 150 * S, 130 * S
    W = LAB + NUM + LW + 30 * S
    H = len(conds) * RH + (len(conds) - 1) * GAP
    im = Image.new("RGB", (W, H), RGB_WHITE)
    d = ImageDraw.Draw(im)
    fb, fm = font(F_BOLD, 27 * S), font(F_MONO, 20 * S)
    for i, c in enumerate(conds):
        y = i * (RH + GAP)
        cy = y + RH / 2
        d.text((0, cy - 17 * S), c["label"], font=fb, fill=RGB_INK)
        d.text((LAB, cy - 13 * S), c.get("count", ""), font=fm, fill=RGB_ACCENT)
        x0 = LAB + NUM
        d.line([x0, cy, x0 + LW, cy], fill=RGB_LINE, width=3 * S)
        for f in c.get("frames", []):
            x = x0 + (f / max(total - 1, 1)) * LW
            x = min(max(x, x0 + 2 * S), x0 + LW - 2 * S)
            d.line([x, y + 6 * S, x, y + RH - 6 * S], fill=RGB_ACCENT, width=5 * S)
    im = im.resize((W // S, H // S), Image.LANCZOS)
    out.parent.mkdir(parents=True, exist_ok=True)
    im.save(out, "PNG")
    return out


# ---------------------------------------------------------------------------
# pptx helpers
# ---------------------------------------------------------------------------
def txt(slide, x, y, w, h, text, *, size=12, bold=False, color=INK, face=SANS,
        align=PP_ALIGN.LEFT, italic=False, spacing=1.15, anchor=MSO_ANCHOR.TOP,
        wrap=True):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    lines = text.split("\n")
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        _rich(p, ln, size, bold, color, face, italic)
    return tb


_TOKEN = re.compile(r"(\*\*[^*]+\*\*|\*[^*\n]+\*|`[^`]+`)")


def _rich(p, line, size, bold, color, face, italic):
    """Render the same **bold** / *italic* / `code` markup the write-ups use, so
    one source sentence can serve both the site and a slide."""
    for tok in _TOKEN.split(line):
        if not tok:
            continue
        text, b, i, fc = tok, bold, italic, face
        if len(tok) > 4 and tok.startswith("**") and tok.endswith("**"):
            text, b = tok[2:-2], True
        elif len(tok) > 2 and tok.startswith("*") and tok.endswith("*"):
            text, i = tok[1:-1], True
        elif len(tok) > 2 and tok.startswith("`") and tok.endswith("`"):
            text, fc = tok[1:-1], MONO
        r = p.add_run()
        r.text = text
        f = r.font
        f.size = Pt(size)
        f.name = fc
        f.bold = b
        f.italic = i
        f.color.rgb = color


def bullets(slide, x, y, w, h, items, *, size=12, color=INK2, gap=8):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = 1.25
        p.space_after = Pt(gap)
        _rich(p, "•  " + it, size, False, color, SANS, False)
    return tb


def rule(slide, x, y, w, color=LINE, h=0.014):
    s = slide.shapes.add_shape(1, Inches(x), Inches(y), Inches(w), Inches(h))  # rect
    s.fill.solid()
    s.fill.fore_color.rgb = color
    s.line.fill.background()
    s.shadow.inherit = False
    return s


def card(slide, x, y, w, h, accent):
    s = slide.shapes.add_shape(5, Inches(x), Inches(y), Inches(w), Inches(h))  # round rect
    s.fill.solid()
    s.fill.fore_color.rgb = PANEL
    s.line.color.rgb = LINE
    s.line.width = Pt(0.75)
    s.shadow.inherit = False
    try:
        s.adjustments[0] = 0.03
    except Exception:
        pass
    rule(slide, x, y, w, accent, h=0.035)
    return s


def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def footer(slide, text, page=None):
    rule(slide, MARGIN, H_IN - 0.62, W_IN - 2 * MARGIN)
    txt(slide, MARGIN, H_IN - 0.50, W_IN - 2 * MARGIN - 0.6, 0.3, text,
        size=8.5, color=FAINT, face=MONO)
    if page:
        txt(slide, W_IN - MARGIN - 0.6, H_IN - 0.50, 0.6, 0.3, page,
            size=8.5, color=FAINT, face=MONO, align=PP_ALIGN.RIGHT)


# ---------------------------------------------------------------------------
# the deck
# ---------------------------------------------------------------------------
class Build:
    def __init__(self, exp_id: str):
        self.exp_id = exp_id
        cdir = ROOT / "content" / exp_id
        self.exp = json.loads((cdir / "exp.json").read_text(encoding="utf-8"))
        self.spec = json.loads((cdir / "slides.json").read_text(encoding="utf-8"))
        raw = self.exp.get("source_dir", "")
        self.src = (ROOT / raw).resolve()
        if not self.src.is_dir():
            sys.exit(f"source_dir not found: {self.src}")
        self.sets = self.exp["media"]["sets"]
        self.seed = (self.exp["media"].get("default_seeds") or ["0"])[0]
        self.conds = self.spec.get("conditions", [])
        self.foot = self.spec.get("footer", "")
        self.page = 0

    # -- media ------------------------------------------------------------
    def video(self, set_name: str, item: str) -> Path | None:
        spec = self.sets.get(set_name)
        if not spec:
            return None
        try:
            name = spec["pattern"].format(id=item, seed=self.seed)
        except Exception:
            return None
        p = self.src / spec.get("dir", "") / name
        return p if p.is_file() else None

    def frames(self, set_name: str, item: str, samples=None) -> list[Path]:
        v = self.video(set_name, item)
        if v is None:
            print(f"    ! missing {set_name}/{item}")
            return []
        return extract(v, samples or SAMPLES, f"{self.exp_id}_{set_name}_{item}")

    def prefetch(self) -> None:
        """One ffmpeg pass per video, in parallel — the slow part of the build."""
        jobs = []
        clips = [s["id"] for s in self.spec["slides"] if s.get("type") == "clip"]
        for c in clips:
            jobs.append(("input", c))
            for cd in self.conds:
                jobs.append((f"vid_{cd['key']}", c))
                jobs.append((f"lift_{cd['key']}", c))
        for s in self.spec["slides"]:
            if s.get("type") == "conditions" and s.get("example"):
                for cd in self.conds:
                    jobs.append((s.get("example_set", "depth_{cond}")
                                 .replace("{cond}", cd["key"]), s["example"]))
            if s.get("type") == "row":
                for iid in s.get("ids", []):
                    for cd in self.conds:
                        jobs.append((s["set"].replace("{cond}", cd["key"]), iid))
        jobs = list(dict.fromkeys(jobs))
        print(f"  extracting frames from {len(jobs)} clips ...")
        with ThreadPoolExecutor(max_workers=8) as ex:
            list(ex.map(lambda j: self.frames(*j), jobs))

    # -- slides -----------------------------------------------------------
    def title_slide(self, prs):
        s = blank(prs)
        it = self.spec.get("intro", {})
        txt(s, MARGIN, 2.05, W_IN - 2 * MARGIN, 0.35,
            self.spec.get("subtitle", ""), size=12, bold=True, color=ACCENT, face=MONO)
        txt(s, MARGIN, 2.45, W_IN - 2 * MARGIN, 1.1, self.spec.get("title", ""),
            size=36, bold=True, color=INK, spacing=1.02)
        if it.get("question"):
            txt(s, MARGIN, 3.75, 9.6, 0.9, it["question"], size=15, color=INK2,
                italic=True, spacing=1.28)
        stats = it.get("stats", [])
        if stats:
            rule(s, MARGIN, 4.95, W_IN - 2 * MARGIN)
            x = MARGIN
            for v, l in stats:
                txt(s, x, 5.15, 2.2, 0.42, str(v), size=22, bold=True, color=INK)
                txt(s, x, 5.60, 2.2, 0.28, l.upper(), size=8.5, color=MUTED, face=MONO)
                x += 2.35
        footer(s, self.foot)

    def what_slide(self, prs):
        s = blank(prs)
        it = self.spec.get("intro", {})
        self.head(s, "What we did", None)
        pipe = pipeline_image(it.get("pipeline", []), it.get("pipeline_side"),
                              CACHE / f"{self.exp_id}_pipeline.png")
        with Image.open(pipe) as im:
            ar = im.width / im.height
        w = W_IN - 2 * MARGIN
        h = w / ar
        s.shapes.add_picture(str(pipe), Inches(MARGIN), Inches(1.15), Inches(w), Inches(h))
        y = 1.15 + h + 0.42
        bullets(s, MARGIN, y, W_IN - 2 * MARGIN, H_IN - y - 0.8,
                it.get("bullets", []), size=12.5, gap=11)
        footer(s, self.foot, self.pagelabel())

    def conditions_slide(self, prs, sl):
        s = blank(prs)
        self.head(s, "The one variable: how many frames carry a pose",
                  "Each tick is a frame the depth control specifies. Everything between "
                  "the ticks is the video model's own work — the control is blank there, "
                  "and the model is told not to attend to it.")
        img = timelines_image(self.conds, self.spec.get("num_frames", 169),
                              CACHE / f"{self.exp_id}_timelines.png")
        with Image.open(img) as im:
            ar = im.width / im.height
        w = W_IN - 2 * MARGIN
        h = min(w / ar, 1.72)
        s.shapes.add_picture(str(img), Inches(MARGIN), Inches(1.72), Inches(w), Inches(h))
        y = 1.72 + h + 0.34
        ex = sl.get("example")
        if ex:
            # Transposed against the clip slides on purpose: conditions across, so
            # the columns sit under the labels above, and time down, so three rows
            # show that both ends are pinned everywhere and only the middle differs.
            tpl = sl.get("example_set", "depth_{cond}")
            times = [0, 56, 168]
            per = {c["key"]: self.frames(tpl.replace("{cond}", c["key"]), ex,
                                         samples=times) for c in self.conds}
            rows = [[(per[c["key"]] or [None] * len(times))[t] for c in self.conds]
                    for t in range(len(times))]
            g = grid_image(rows, CACHE / f"{self.exp_id}_condex.jpg", tile_w=320)
            if g:
                txt(s, MARGIN, y, 8.0, 0.26,
                    f"What that looks like \u2014 the control track for "
                    f"\u201c{ex.replace('_', ' ')}\u201d",
                    size=11, bold=True, color=INK)
                self.labelled_grid(s, g, [f"f{t}" for t in times],
                                   y + 0.62, MARGIN, W_IN - 2 * MARGIN,
                                   H_IN - (y + 0.62) - 0.80, label_w=0.70,
                                   col_labels=[c["label"] for c in self.conds],
                                   col_y=y + 0.36, label_face=MONO)
        footer(s, self.foot, self.pagelabel())

    def clip_slide(self, prs, sl):
        cid = sl["id"]
        info = self.spec.get("clips", {}).get(cid, {})
        s = blank(prs)
        nice = cid.replace("_", " ")
        txt(s, MARGIN, 0.34, 4.2, 0.42, nice, size=25, bold=True, color=INK)
        txt(s, MARGIN + 1.05 + 0.15 * len(nice), 0.44, 7.0, 0.32,
            f"“{info.get('prompt', '')}”", size=13, color=MUTED, italic=True)
        txt(s, MARGIN, 0.86, 9.55, 0.62, info.get("headline", ""), size=11.5,
            color=INK2, spacing=1.28)

        # input motion, top right — what went in
        fin = self.frames("input", cid, samples=[0, 84, 168])
        gin = grid_image([fin], CACHE / f"{self.exp_id}_in_{cid}.jpg", tile_w=200)
        if gin:
            with Image.open(gin) as im:
                ar = im.width / im.height
            w = 2.55
            s.shapes.add_picture(str(gin), Inches(W_IN - MARGIN - w), Inches(0.36),
                                 Inches(w), Inches(w / ar))
            txt(s, W_IN - MARGIN - w, 0.36 + w / ar + 0.06, w, 0.24,
                "INPUT MOTION · WHAT WENT IN", size=7.5, color=FAINT, face=MONO,
                align=PP_ALIGN.RIGHT)

        # two halves: what the model made | what came back
        gy = 2.02
        gh = H_IN - gy - 0.78
        lab_w = 1.02
        half = (W_IN - 2 * MARGIN - lab_w - 0.34) / 2
        for k, (set_tpl, title) in enumerate([
                ("vid_{cond}", "What the model made — photoreal output"),
                ("lift_{cond}", "What came back as motion — recovered SMPL")]):
            rows = [self.frames(set_tpl.replace("{cond}", c["key"]), cid)
                    for c in self.conds]
            g = grid_image(rows, CACHE / f"{self.exp_id}_{set_tpl[:3]}_{cid}.jpg")
            x = MARGIN + lab_w + k * (half + 0.34)
            txt(s, x, 1.56, half, 0.26, title, size=10.5, bold=True, color=INK)
            for ci, n in enumerate(SAMPLES):
                cw = half / len(SAMPLES)
                txt(s, x + ci * cw, 1.83, cw, 0.2, f"f{n}", size=7.5, color=FAINT,
                    face=MONO, align=PP_ALIGN.CENTER)
            if g:
                s.shapes.add_picture(str(g), Inches(x), Inches(gy), Inches(half),
                                     Inches(gh))
        # row labels: condition + its joint error
        errs = info.get("errors") or []
        best = min(errs) if errs else None
        rh = gh / len(self.conds)
        for i, c in enumerate(self.conds):
            yy = gy + i * rh
            txt(s, MARGIN, yy + rh / 2 - 0.20, lab_w - 0.08, 0.24, c["label"],
                size=11, bold=True, color=INK)
            if i < len(errs):
                col = KEY if errs[i] == best else MUTED
                txt(s, MARGIN, yy + rh / 2 + 0.02, lab_w - 0.08, 0.22,
                    f"{errs[i]} mm", size=9, bold=errs[i] == best, color=col, face=MONO)
        chips = " · ".join(x for x in (info.get("action"), info.get("breaks")) if x)
        footer(s, f"{chips}    |    joint error vs. the input motion, root-relative"
               if chips else self.foot, self.pagelabel())

    def row_slide(self, prs, sl):
        s = blank(prs)
        self.head(s, sl.get("title", ""), sl.get("headline", ""))
        ids = sl.get("ids", [])
        rows = []
        for iid in ids:
            rows.append([(self.frames(sl["set"].replace("{cond}", c["key"]), iid) or [None])[2]
                         for c in self.conds])
        g = grid_image(rows, CACHE / f"{self.exp_id}_row.jpg", tile_w=380)
        gy, lab_w = 2.16, 1.02
        gh = H_IN - gy - 0.82
        if g:
            self.labelled_grid(s, g, [i.replace("_", " ") for i in ids], gy, MARGIN,
                               W_IN - 2 * MARGIN, gh, label_w=lab_w,
                               col_labels=[c["label"] for c in self.conds],
                               col_y=1.86)
        if sl.get("note"):
            txt(s, MARGIN, H_IN - 1.02, W_IN - 2 * MARGIN, 0.30, sl["note"],
                size=9.5, color=MUTED)
        footer(s, self.foot, self.pagelabel())

    def closing_slide(self, prs):
        cl = self.spec.get("closing", {})
        s = blank(prs)
        self.head(s, cl.get("title", "What the sweep shows"), cl.get("headline", ""))
        cols = cl.get("columns", [])
        if cols:
            gap = 0.34
            w = (W_IN - 2 * MARGIN - gap * (len(cols) - 1)) / len(cols)
            top, h = 2.10, 3.55
            for i, c in enumerate(cols):
                x = MARGIN + i * (w + gap)
                accent = {"good": KEY, "bad": HEAVY}.get(c.get("kind"), ACCENT)
                card(s, x, top, w, h, accent)
                txt(s, x + 0.24, top + 0.28, w - 0.48, 0.5, c.get("title", ""),
                    size=12.5, bold=True, color=INK)
                bullets(s, x + 0.24, top + 0.86, w - 0.48, h - 1.1,
                        c.get("items", []), size=9.5, gap=6)
        if cl.get("punchline"):
            txt(s, MARGIN, 6.14, W_IN - 2 * MARGIN, 0.5, cl["punchline"], size=14,
                color=INK, align=PP_ALIGN.CENTER, spacing=1.2)
        footer(s, self.foot, self.pagelabel())

    # -- shared -----------------------------------------------------------
    def head(self, slide, title, sub):
        txt(slide, MARGIN, 0.34, W_IN - 2 * MARGIN, 0.5, title, size=23, bold=True,
            color=INK)
        if sub:
            txt(slide, MARGIN, 0.92, W_IN - 2 * MARGIN - 1.2, 0.62, sub, size=11.5,
                color=INK2, spacing=1.28)

    def labelled_grid(self, slide, img, row_labels, y, x, w, h, *, label_w=1.0,
                      col_labels=None, col_y=None, label_face=SANS):
        with Image.open(img) as im:
            ar = im.width / im.height
        gw = w - label_w
        gh = min(gw / ar, h)
        gw = gh * ar
        gx = x + label_w
        slide.shapes.add_picture(str(img), Inches(gx), Inches(y), Inches(gw), Inches(gh))
        rh = gh / max(len(row_labels), 1)
        for i, lab in enumerate(row_labels):
            txt(slide, x, y + i * rh + rh / 2 - 0.11, label_w - 0.08, 0.24, lab,
                size=10.5, bold=True, color=INK, face=label_face)
        if col_labels and col_y is not None:
            cw = gw / len(col_labels)
            for i, lab in enumerate(col_labels):
                txt(slide, gx + i * cw, col_y, cw, 0.24, lab, size=10.5, bold=True,
                    color=INK, align=PP_ALIGN.CENTER)

    def pagelabel(self) -> str:
        return str(self.page)

    # -- run --------------------------------------------------------------
    def run(self) -> Path:
        self.prefetch()
        prs = Presentation()
        prs.slide_width, prs.slide_height = Inches(W_IN), Inches(H_IN)
        print("  laying out slides ...")
        self.title_slide(prs)
        self.page = 1
        self.what_slide(prs)
        for sl in self.spec.get("slides", []):
            t = sl.get("type")
            if t == "intro":
                continue
            self.page += 1
            if t == "conditions":
                self.conditions_slide(prs, sl)
            elif t == "clip":
                self.clip_slide(prs, sl)
            elif t == "row":
                self.row_slide(prs, sl)
            elif t == "closing":
                self.closing_slide(prs)
            else:
                self.page -= 1
        out = OUT_DIR / f"{self.exp_id}-slides.pptx"
        prs.save(str(out))
        print(f"\n  {len(prs.slides._sldIdLst)} slides")
        print(f"  {out.name}  ({out.stat().st_size / 1e6:.1f} MB)  ->  {out.parent}")
        print("  light theme, pictures and text boxes only, no embedded media")
        return out


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("exp", nargs="?", default="exp6", help="experiment id, e.g. exp6")
    a = ap.parse_args()
    Build(a.exp).run()
    return 0


if __name__ == "__main__":
    sys.exit(main())
